"""Export content/tools.yaml to an Excel register you can edit and share.   python export_tools.py   (needs openpyxl)"""
import yaml
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).parent
cfg = yaml.safe_load((ROOT / "content" / "tools.yaml").read_text(encoding="utf-8"))
cols = cfg["columns"]; keys = list(cols)
wb = Workbook(); ws = wb.active; ws.title = "Tools"
ws.append(["Tool", "Owner / origin", *[cols[k] for k in keys], "Source", "In diagram"])
for t in cfg["tools"]:
    ws.append([t["name"], t["owner"], *[t.get(k) or "?" for k in keys], t.get("source", ""), "yes" if t.get("diagram") else ""])
for c in ws[1]:
    c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="93002D")
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name="Arial", size=10, color="B00052" if c.value == "?" else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")
for col, w in zip("ABCDEFGHIJ", [34, 30, 34, 34, 30, 26, 30, 26, 40, 10]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "B2"
ws2 = wb.create_sheet("Shared core layers"); ws2.append(["Layer", "What it covers"])
for l in cfg["core_layers"]:
    ws2.append([l["title"], l["text"]])
ws2.column_dimensions["A"].width = 36; ws2.column_dimensions["B"].width = 110
out = ROOT / "tools_register.xlsx"; wb.save(out); print("Wrote", out)
